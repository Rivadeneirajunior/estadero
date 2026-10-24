"""Modelo del caso base: estadero abierto entre semana en 2027.
Lee datos/supuestos.xlsx (nada escrito a mano aquí), calcula,
escribe salidas/modelo.xlsx (con fórmulas) y los resultados dentro de salidas/index.html."""
import calendar, datetime, json, math, re
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

RAIZ = Path(__file__).resolve().parent.parent
ANIO = 2027
MESES = ["Enero","Febrero","Marzo","Abril","Mayo","Junio","Julio","Agosto","Septiembre","Octubre","Noviembre","Diciembre"]

# 1. Leer supuestos (solo lectura)
hoja = load_workbook(RAIZ / "datos" / "supuestos.xlsx", data_only=True)["Supuestos"]
filas = [f for f in hoja.iter_rows(min_row=2, values_only=True) if f[0]]
S = {f[0]: f for f in filas}
poblacion = S["Personas en el área de influencia"][1]
captura = S["Captura por día abierto"][1]
precio = S["Precio por persona"][1]
costo = S["Costo por cliente"][1]
fijos = S["Costos fijos adicionales al mes"][1]
patron = str(S["Patrón de días abiertos entre semana"][1]).zfill(7)

def redondear(x):  # mitades hacia arriba
    return math.floor(x + 0.5)

def dias_abiertos(mes):
    n = calendar.monthrange(ANIO, mes)[1]
    return sum(1 for d in range(1, n + 1)
               if patron[datetime.date(ANIO, mes, d).weekday()] == "0")

# 2. Cálculo mes a mes
meses, acum, minimo = [], 0, 0
for m in range(1, 13):
    dias = dias_abiertos(m)
    clientes = redondear(dias * poblacion * captura)
    ing, cv = clientes * precio, clientes * costo
    util = ing - cv - fijos
    acum += util
    minimo = min(minimo, acum)
    meses.append(dict(mes=MESES[m - 1], dias=dias, clientes=clientes, ingresos=ing, cv=cv, cf=fijos, utilidad=util))
tot = {k: sum(x[k] for x in meses) for k in ("dias", "clientes", "ingresos", "cv", "cf", "utilidad")}
eq_exacto = fijos / (precio - costo)
eq = math.ceil(eq_exacto)
perdidas = sum(1 for x in meses if x["utilidad"] < 0)
reserva = -minimo if minimo < 0 else 0

# 2b. Prueba del pronóstico (instrucción 2)
import csv
hist = {int(r["anio"]): int(r["clientes"]) for r in csv.DictReader(open(RAIZ / "datos" / "limpios" / "historia-clientes.csv"))}
a, b, c = sorted(hist)[-3:]
real = hist[c]
pron = [dict(metodo=f"Seguir la tendencia ({c} crece como de {a} a {b})", pred=round(hist[b] * hist[b] / hist[a])),
        dict(metodo=f"Repetir {b}", pred=hist[b])]
for x in pron: x["real"] = real; x["error"] = round((x["pred"] - real) / real * 100, 1)
# 2c. Escenarios (instrucción 3): solo cambia la demanda
def simular(var):
    ut, ac, mn, cls = [], 0, 0, []
    for m in meses:
        cli = redondear(m["dias"] * poblacion * captura * (1 + var))
        u = cli * (precio - costo) - fijos
        ac += u; mn = min(mn, ac); ut.append(u); cls.append(cli)
    return dict(var=var, utilidades=ut, clientes=sum(cls), utilidad=sum(ut), perdidas=sum(1 for u in ut if u < 0), reserva=-mn if mn < 0 else 0)
esc = {n: simular(S[f"Variación de demanda, escenario {n}"][1]) for n in ("pesimista", "base", "optimista")}
assert esc["base"]["utilidad"] == tot["utilidad"]

# 2d. Sensibilidad (instrucción 4): un supuesto a la vez, dentro de su rango
def util(pob=poblacion, cap=captura, pr=precio, co=costo, fj=fijos):
    return sum(redondear(m["dias"] * pob * cap) * (pr - co) - fj for m in meses)
assert util() == tot["utilidad"]
base_par = dict(pob=poblacion, cap=captura, pr=precio, co=costo, fj=fijos)
torn = []
for nombre, k in [("Captura por día abierto", "cap"), ("Personas en el área de influencia", "pob"), ("Precio por persona", "pr"),
                  ("Costo por cliente", "co"), ("Costos fijos adicionales al mes", "fj")]:
    r = S[nombre][5]
    d = [util(**{k: base_par[k] * (1 + sg * r)}) - tot["utilidad"] for sg in (-1, 1)]
    torn.append(dict(nombre=nombre, rango=round(r * 100), mal=round(min(d)), bien=round(max(d))))
torn.sort(key=lambda x: x["bien"] - x["mal"], reverse=True)
cap_eq = (fijos * 12 / (precio - costo)) / (tot["dias"] * poblacion)
par = dict(base_par, src={k: S[n][4] for k, n in [("pob", "Personas en el área de influencia"), ("cap", "Captura por día abierto"), ("pr", "Precio por persona"), ("co", "Costo por cliente"), ("fj", "Costos fijos adicionales al mes")]})

# 2e. Memo de decisión (instrucción 6): solo cifras que están en modelo.xlsx
def P(n): return ("-" if n < 0 else "") + "$ " + f"{abs(n):,.0f}".replace(",", ".")
def pc(x, d=2): return f"{x * 100:.{d}f} %".replace(".", ",")
e = esc
memo = dict(
  reco="Abrir entre semana (miércoles y jueves), pero con un piloto de tres meses antes de comprometer un año de costos fijos adicionales.",
  esc=f"Si la demanda es {pc(abs(e['pesimista']['var']), 0)} menor, la utilidad del año sería {P(e['pesimista']['utilidad'])} con {e['pesimista']['perdidas']} meses en pérdida; "
      f"en el caso base, {P(e['base']['utilidad'])} con {e['base']['perdidas']} meses en pérdida; si es {pc(e['optimista']['var'], 0)} mayor, {P(e['optimista']['utilidad'])} con {e['optimista']['perdidas']} meses en pérdida.",
  riesgo=f"El supuesto más riesgoso es la captura (rango ± {round(S['Captura por día abierto'][5] * 100)} %, el más amplio). Si menos de {pc(cap_eq, 3)} de las personas del área visita en un día abierto (el caso base supone {pc(captura)}), el año no cubre los costos adicionales.",
  comprobar="Medir durante un piloto de tres meses, o con una preventa o encuesta, cuántas personas llegan por día abierto y compararlo con la captura de equilibrio.",
  reserva=f"Reserva necesaria: {P(e['pesimista']['reserva'])} (escenario pesimista); el caso base no necesita reserva. No incluye inversión inicial.",
  datos="Lo que dicen los datos: todavía no hay datos medidos. La historia de clientes y todos los valores de este ejercicio son inventados.",
  sup=f"Lo que suponemos: {str(poblacion)[:-3]}.{str(poblacion)[-3:]} personas en el área, captura de {pc(captura)}, precio de {P(precio)}, costo de {P(costo)} por cliente y {P(fijos)} al mes de costos fijos adicionales. "
      f"Si se cumplen, tendríamos {tot['clientes']:,} clientes y una utilidad de {P(tot['utilidad'])} en {ANIO}.".replace(f"{tot['clientes']:,}", f"{tot['clientes']:,}".replace(",", ".")))
(RAIZ / "salidas").mkdir(exist_ok=True)
(RAIZ / "salidas" / "MEMO.md").write_text("# Decisión: ¿abrir el estadero entre semana en 2027?\n\n**Recomendación:** " + memo["reco"] + "\n\n**Escenarios:** " + memo["esc"] + "\n\n**Supuesto más riesgoso:** " + memo["riesgo"] + "\n\n**Cómo comprobarlo:** " + memo["comprobar"] + "\n\n**" + memo["reserva"].replace(":", ":**", 1) + "\n\n---\n" + memo["datos"] + "\n\n" + memo["sup"] + "\n", encoding="utf-8")

# 3. Excel con fórmulas que apuntan a la hoja Supuestos
wb = Workbook()
ws = wb.active; ws.title = "Supuestos"
ws.append(["Supuesto", "Valor", "Unidad", "Tipo", "De dónde sale", "Rango posible"])
for f in filas: ws.append(list(f))
for r in range(2, ws.max_row + 1):
    ws.cell(r, 2).number_format = '0.00%' if ws.cell(r, 1).value.startswith("Captura") else ('0%' if ws.cell(r, 1).value.startswith("Variación") else '#,##0')
    ws.cell(r, 6).number_format = '"± "0%'
ws["B7"].number_format = "@"
cs = wb.create_sheet("Cálculo")
cs.append(["Mes", "Días abiertos", "Clientes", "Ingresos", "Costos variables", "Costos fijos", "Utilidad", "Utilidad acumulada"])
for m in range(1, 13):
    r = m + 1
    cs.append([MESES[m - 1],
        f'=NETWORKDAYS.INTL(DATE({ANIO},{m},1),EOMONTH(DATE({ANIO},{m},1),0),Supuestos!$B$7)',
        f'=ROUND(B{r}*Supuestos!$B$2*Supuestos!$B$3,0)',
        f'=C{r}*Supuestos!$B$4', f'=C{r}*Supuestos!$B$5', '=Supuestos!$B$6',
        f'=D{r}-E{r}-F{r}', f'=SUM($G$2:G{r})'])
cs.append(["Total"] + [f"=SUM({c}2:{c}13)" for c in "BCDEFG"])
cs["A16"], cs["B16"] = "Punto de equilibrio (clientes al mes)", "=Supuestos!B6/(Supuestos!B4-Supuestos!B5)"
cs["A17"], cs["B17"] = "Meses con pérdida", '=COUNTIF(G2:G13,"<0")'
cs["A18"], cs["B18"] = "Reserva necesaria", "=MAX(0,-MIN(H2:H13))"
cs["A19"], cs["B19"] = "Captura de equilibrio (% de las personas del área, por día abierto)", "=(Supuestos!B6*12/(Supuestos!B4-Supuestos!B5))/(B14*Supuestos!B2)"
for ws_ in (ws, cs):
    for row in ws_.iter_rows():
        for c in row: c.font = Font(name="Arial")
    for c in ws_[1]: c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F6B3F")
for row in cs.iter_rows(min_row=2, min_col=2):
    for c in row: c.number_format = '#,##0;-#,##0;0'
cs["B16"].number_format = '#,##0.0'; cs["B19"].number_format = '0.000%'
for col, w in zip("ABCDEFGH", (38, 14, 12, 16, 18, 16, 16, 20)): cs.column_dimensions[col].width = w
for col, w in zip("ABCDEF", (38, 14, 40, 10, 80, 14)): ws.column_dimensions[col].width = w
es = wb.create_sheet("Escenarios")
es.append(["Mes", "Pesimista", "Base", "Optimista", "Acum. pesimista", "Acum. base", "Acum. optimista"])
for m in range(1, 13):
    r = m + 1
    es.append([MESES[m - 1]] + [f"=ROUND('Cálculo'!$B{r}*Supuestos!$B$2*Supuestos!$B$3*(1+{c}$17),0)*(Supuestos!$B$4-Supuestos!$B$5)-Supuestos!$B$6" for c in "BCD"]
              + [f"=SUM({c}$2:{c}{r})" for c in "BCD"])
es.append(["Utilidad del año"] + [f"=SUM({c}2:{c}13)" for c in "BCD"])
es.append(["Meses con pérdida"] + [f'=COUNTIF({c}2:{c}13,"<0")' for c in "BCD"])
es.append(["Reserva necesaria"] + [f"=MAX(0,-MIN({c}2:{c}13))" for c in "EFG"])
es.append(["Variación de demanda", "=Supuestos!B8", "=Supuestos!B9", "=Supuestos!B10"])
for row in es.iter_rows():
    for c in row: c.font = Font(name="Arial"); c.number_format = '#,##0;-#,##0;0'
for c in es[1]: c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F6B3F")
for c in es[17][1:]: c.number_format = '0%'
es.column_dimensions["A"].width = 24
for col in "BCDEFG": es.column_dimensions[col].width = 18
(RAIZ / "salidas").mkdir(exist_ok=True)
wb.save(RAIZ / "salidas" / "modelo.xlsx")

# 4. Escribir resultados dentro del dashboard
datos = dict(torn=torn, par=par, memo=memo, pron=pron, esc=esc, meses=meses, tot=tot, eq=eq, perdidas=perdidas, reserva=reserva,
             sup=[[f[0], f"{f[1]:,}".replace(",", ".") if isinstance(f[1], (int, float)) and f[0] != "Captura por día abierto" else (f"{f[1]*100:.2f} %".replace(".", ",") if isinstance(f[1], float) else f[1]), f[4], round(f[5] * 100)] for f in filas])
html = RAIZ / "salidas" / "index.html"
txt = html.read_text(encoding="utf-8")
txt = re.sub(r"/\*DATA\*/.*?/\*FIN\*/", lambda _: "/*DATA*/" + json.dumps(datos, ensure_ascii=False) + "/*FIN*/", txt, flags=re.S)
html.write_text(txt, encoding="utf-8")

# 5. Cifras de control en consola
print("Mes | Días | Clientes | Ingresos | Costos var. | Costos fijos | Utilidad")
for x in meses: print(x["mes"], x["dias"], x["clientes"], x["ingresos"], x["cv"], x["cf"], x["utilidad"])
print("TOTAL", tot); print("TORNADO", torn); print("CAP_EQ", cap_eq)
print("PRONÓSTICO", pron)
for n, e in esc.items(): print("ESCENARIO", n, e["utilidad"], e["perdidas"], e["reserva"])
print(f"Punto de equilibrio: {eq_exacto:.1f} (≈ {eq}) clientes al mes | Meses con pérdida: {perdidas} | Reserva: {reserva}")
